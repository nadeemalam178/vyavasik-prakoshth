import openpyxl
import json
import re
from collections import defaultdict, Counter

def process_all_data(excel_path='data.xlsx'):
    wb = openpyxl.load_workbook(excel_path, data_only=True)

    # -------------------------------------------------------------
    # 1. Parse Campaign Committe
    # -------------------------------------------------------------
    ws_cc = wb['Campaign Committe']
    campaign_committee = []
    current_role = 'Members'
    current_zone = ''

    for r in range(1, ws_cc.max_row + 1):
        val_a = ws_cc.cell(r, 1).value
        val_a_str = str(val_a).strip() if val_a is not None else ''
        
        lower_a = val_a_str.lower()
        if any(kw in lower_a for kw in ['sanyojak', 'members', 'committee']):
            if 'kshetriya' in lower_a:
                current_role = 'Kshetriya Sanyojak'
            elif 'sah' in lower_a:
                current_role = 'Sah Sanyojak'
            elif 'sanyojak' in lower_a:
                current_role = 'Sanyojak'
            else:
                current_role = val_a_str
            continue
        if lower_a in ['sr no', 'sr. no', 'sr_no', 's no']:
            continue
        
        sr = ws_cc.cell(r, 1).value
        zone = ws_cc.cell(r, 2).value
        district = ws_cc.cell(r, 3).value
        name = ws_cc.cell(r, 4).value
        number = ws_cc.cell(r, 5).value
        age = ws_cc.cell(r, 6).value
        category = ws_cc.cell(r, 7).value
        caste = ws_cc.cell(r, 8).value
        profile = ws_cc.cell(r, 9).value
        
        if zone:
            current_zone = str(zone).strip()
            
        if name or district or (sr and isinstance(sr, (int, float))):
            phone_clean = ''
            if number:
                if isinstance(number, (int, float)):
                    phone_clean = str(int(number))
                else:
                    phone_clean = str(number).strip()
                    if phone_clean.endswith('.0'): phone_clean = phone_clean[:-2]
            
            age_clean = None
            if age:
                try: age_clean = int(float(age))
                except: pass
                    
            cat_clean = str(category).strip() if category else ''
            if cat_clean.lower() in ['gen', 'general']: cat_clean = 'General'
                
            caste_clean = str(caste).strip() if caste else ''
            if caste_clean.lower() == 'bania': caste_clean = 'Baniya'
            elif 'kalwar' in caste_clean.lower(): caste_clean = 'Kalwar'
            elif 'brahmin' in caste_clean.lower() or 'bhraman' in caste_clean.lower(): caste_clean = 'Brahmin'
                
            name_clean = str(name).strip() if name else '(Vacant / To be Appointed)'
            is_vacant = not bool(name and str(name).strip())
            
            # Extract profile tags
            p_lower = str(profile).lower() if profile else ''
            parties = []
            if 'bjp' in p_lower: parties.append('Former BJP')
            if 'jdu' in p_lower or 'jd(u)' in p_lower: parties.append('Former JDU')
            if 'congress' in p_lower or 'inc' in p_lower: parties.append('Former Congress')
            if 'rjd' in p_lower: parties.append('Former RJD')
            if 'ljp' in p_lower: parties.append('Former LJP')
            
            sectors = []
            if any(k in p_lower for k in ['hotel', 'resort', 'restaurant']): sectors.append('Hospitality')
            if any(k in p_lower for k in ['poultry', 'hatchery', 'egg', 'cattle']): sectors.append('Poultry & Agro')
            if any(k in p_lower for k in ['jewell', 'swarn', 'gold']): sectors.append('Jewellery')
            if any(k in p_lower for k in ['nursing', 'doctor', 'hospital', 'medicine', 'drug', 'chemist']): sectors.append('Healthcare & Pharma')
            if any(k in p_lower for k in ['auto', 'tvs', 'tyre', 'electric', 'ev']): sectors.append('Automobile & EV')
            if any(k in p_lower for k in ['contractor', 'construction', 'building', 'cement', 'real estate', 'mining', 'mines']): sectors.append('Construction & Real Estate')
            if any(k in p_lower for k in ['college', 'school', 'vidhyalay', 'education']): sectors.append('Education')
            if any(k in p_lower for k in ['petrol', 'gas agency']): sectors.append('Petroleum & Energy')
            if any(k in p_lower for k in ['transport', 'truck', 'bus']): sectors.append('Logistics & Transport')
            if any(k in p_lower for k in ['advocate', 'lawyer', 'llb']): sectors.append('Legal')
            if any(k in p_lower for k in ['mukhiya', 'sarpanch', 'ward parshad', 'pacs', 'parshad', 'nagar parishad', 'bdc']): sectors.append('Grassroots Elected')
            
            campaign_committee.append({
                'sr': int(sr) if isinstance(sr, (int, float)) else len(campaign_committee) + 1,
                'role': current_role,
                'role_tier': 1 if current_role == 'Sanyojak' else (2 if current_role == 'Sah Sanyojak' else 3),
                'zone': current_zone,
                'district': str(district).strip() if district else '',
                'name': name_clean,
                'is_vacant': is_vacant,
                'number': phone_clean,
                'age': age_clean,
                'category': cat_clean,
                'caste': caste_clean,
                'profile': str(profile).strip() if profile else '',
                'parties': parties,
                'sectors': sectors
            })

    # -------------------------------------------------------------
    # 2. Parse GAP Report
    # -------------------------------------------------------------
    ws_gap = wb['GAP Report']
    gap_by_district = {}
    for r in range(4, ws_gap.max_row + 1):
        d_name = ws_gap.cell(r, 2).value
        if not d_name or str(d_name).strip().lower() == 'total':
            break
        target = ws_gap.cell(r, 3).value
        achieved = ws_gap.cell(r, 4).value
        gap = ws_gap.cell(r, 5).value
        poc = ws_gap.cell(r, 6).value
        clean_name = str(d_name).strip().replace('\n', '').strip()
        gap_by_district[clean_name.lower()] = {
            'district': clean_name,
            'target': int(target) if target else 0,
            'achieved': int(achieved) if achieved else 0,
            'gap': int(gap) if gap else 0,
            'poc': str(poc).strip() if poc else ''
        }

    # -------------------------------------------------------------
    # 3. Parse Vyvsaik SWG+DPGS
    # -------------------------------------------------------------
    ws_swg = wb['Vyvsaik SWG+DPGS']
    swg_members = []
    rows_swg = list(ws_swg.iter_rows(values_only=True))

    rank_map = {
        'State President': 1,
        'State Gen Sec': 2,
        'State VP': 3,
        'Zonal': 4,
        'District President': 5,
        'District Gen Sec': 6,
        'State Core Committee': 7
    }

    for idx, r in enumerate(rows_swg[1:], start=1):
        if not any(r): continue
        
        row_vals = (list(r) + [None] * 12)[:11]
        sr, zone, dist, post, _col_rank, name, num, age, cat, caste, prof = row_vals
        
        post_str = str(post).strip() if post else ''
        if post_str not in rank_map:
            continue
        if str(sr).strip().lower() in ['sr no', 'sr. no', 'sr_no', 's no']:
            continue
        
        zone_str = str(zone).strip() if zone else ''
        dist_str = str(dist).strip() if dist else ''
        post_str = str(post).strip() if post else ''
        name_str = str(name).strip() if name else ''
        is_vacant = not bool(name_str)
        
        phone_clean = ''
        if num:
            if isinstance(num, (int, float)): phone_clean = str(int(num))
            else:
                phone_clean = str(num).strip()
                if phone_clean.endswith('.0'): phone_clean = phone_clean[:-2]
                
        age_clean = None
        if age:
            try: age_clean = int(float(age))
            except: pass
                
        cat_clean = str(cat).strip() if cat else ''
        if cat_clean.lower() in ['gen', 'general']: cat_clean = 'General'
        elif cat_clean.lower() == 'baudhist': cat_clean = 'Minority (Buddhist)'
        elif cat_clean.lower() == 'minority': cat_clean = 'Minority'
            
        caste_clean = str(caste).strip() if caste else ''
        if caste_clean.lower() in ['bania', 'baniya']: caste_clean = 'Baniya'
        elif 'kalwar' in caste_clean.lower(): caste_clean = 'Kalwar'
        elif 'brahmin' in caste_clean.lower() or 'bhraman' in caste_clean.lower(): caste_clean = 'Brahmin'
            
        rank = rank_map.get(post_str, 99)
        
        # Profile Intelligence Extraction
        p_lower = str(prof).lower() if prof else ''
        parties = []
        if 'bjp' in p_lower: parties.append('Former BJP')
        if 'jdu' in p_lower or 'jd(u)' in p_lower: parties.append('Former JDU')
        if 'congress' in p_lower or 'inc' in p_lower: parties.append('Former Congress')
        if 'rjd' in p_lower: parties.append('Former RJD')
        if 'ljp' in p_lower: parties.append('Former LJP')
        
        sectors = []
        if any(k in p_lower for k in ['hotel', 'resort', 'restaurant']): sectors.append('Hospitality')
        if any(k in p_lower for k in ['poultry', 'hatchery', 'egg', 'cattle']): sectors.append('Poultry & Agro')
        if any(k in p_lower for k in ['jewell', 'swarn', 'gold']): sectors.append('Jewellery')
        if any(k in p_lower for k in ['nursing', 'doctor', 'hospital', 'medicine', 'drug', 'chemist']): sectors.append('Healthcare & Pharma')
        if any(k in p_lower for k in ['auto', 'tvs', 'tyre', 'electric', 'ev']): sectors.append('Automobile & EV')
        if any(k in p_lower for k in ['contractor', 'construction', 'building', 'cement', 'real estate', 'mining', 'mines']): sectors.append('Construction & Real Estate')
        if any(k in p_lower for k in ['college', 'school', 'vidhyalay', 'education']): sectors.append('Education')
        if any(k in p_lower for k in ['petrol', 'gas agency']): sectors.append('Petroleum & Energy')
        if any(k in p_lower for k in ['transport', 'truck', 'bus']): sectors.append('Logistics & Transport')
        if any(k in p_lower for k in ['advocate', 'lawyer', 'llb']): sectors.append('Legal')
        if any(k in p_lower for k in ['wholesale', 'trading', 'shop', 'retail', 'dealer', 'vyapar']): sectors.append('Trading & Retail')
        if any(k in p_lower for k in ['mukhiya', 'sarpanch', 'ward parshad', 'pacs', 'parshad', 'nagar parishad', 'bdc']): sectors.append('Grassroots Elected')
        
        swg_members.append({
            'id': idx,
            'sr_no': sr,
            'zone': zone_str,
            'district': dist_str,
            'post': post_str,
            'post_rank': rank,
            'name': name_str if not is_vacant else '(Vacant)',
            'is_vacant': is_vacant,
            'number': phone_clean,
            'age': age_clean,
            'category': cat_clean,
            'caste': caste_clean,
            'profile': str(prof).strip() if prof else '',
            'parties': parties,
            'sectors': sectors
        })

    # -------------------------------------------------------------
    # 4. Parse Tirhut Calling Status
    # -------------------------------------------------------------
    ws_tirhut = wb['Tirhut Calling Status']
    tirhut_calls = []
    calling_status_counts = Counter()

    for r in range(2, ws_tirhut.max_row + 1):
        sr = ws_tirhut.cell(r, 1).value
        if not sr: continue
        d_name = ws_tirhut.cell(r, 2).value
        post = ws_tirhut.cell(r, 4).value
        name = ws_tirhut.cell(r, 5).value
        num = ws_tirhut.cell(r, 6).value
        poc = ws_tirhut.cell(r, 7).value
        status = ws_tirhut.cell(r, 8).value
        
        phone_str = ''
        if num:
            if isinstance(num, (int, float)): phone_str = str(int(num))
            else:
                phone_str = str(num).strip()
                if phone_str.endswith('.0'): phone_str = phone_str[:-2]
                
        status_clean = str(status).strip() if status else 'Pending'
        calling_status_counts[status_clean] += 1
        
        tirhut_calls.append({
            'sr': sr,
            'district': str(d_name).strip() if d_name else '',
            'post': str(post).strip() if post else '',
            'name': str(name).strip() if name else '',
            'phone': phone_str,
            'poc': str(poc).strip() if poc else '',
            'status': status_clean
        })

    # -------------------------------------------------------------
    # 5. Zone & District Aggregations
    # -------------------------------------------------------------
    zones_map = defaultdict(lambda: {
        'name': '',
        'districts': set(),
        'total_posts': 0,
        'filled_posts': 0,
        'vacant_posts': 0,
        'zonal_lead': None
    })

    districts_map = defaultdict(lambda: {
        'name': '',
        'zone': '',
        'total_posts': 0,
        'filled_posts': 0,
        'vacant_posts': 0,
        'district_president': None,
        'district_gen_sec': None,
        'core_committee_count': 0,
        'gap_target': 7,
        'gap_achieved': 0,
        'gap_remaining': 0,
        'poc': ''
    })

    for m in swg_members:
        z = m['zone']
        d = m['district']
        
        if z:
            zones_map[z]['name'] = z
            if d: zones_map[z]['districts'].add(d)
            zones_map[z]['total_posts'] += 1
            if m['is_vacant']: zones_map[z]['vacant_posts'] += 1
            else:
                zones_map[z]['filled_posts'] += 1
                if m['post'] == 'Zonal': zones_map[z]['zonal_lead'] = m['name']

        if d:
            districts_map[d]['name'] = d
            districts_map[d]['zone'] = z
            districts_map[d]['total_posts'] += 1
            if m['is_vacant']: districts_map[d]['vacant_posts'] += 1
            else:
                districts_map[d]['filled_posts'] += 1
                if m['post'] == 'District President': districts_map[d]['district_president'] = m['name']
                elif m['post'] == 'District Gen Sec': districts_map[d]['district_gen_sec'] = m['name']
                elif m['post'] == 'State Core Committee': districts_map[d]['core_committee_count'] += 1

    for d_name, d_obj in districts_map.items():
        match = None
        key = d_name.lower()
        if key in gap_by_district:
            match = gap_by_district[key]
        else:
            for k, v in gap_by_district.items():
                if k in key or key in k:
                    match = v
                    break
        if match:
            d_obj['gap_target'] = match['target']
            d_obj['gap_achieved'] = match['achieved']
            d_obj['gap_remaining'] = match['gap']
            d_obj['poc'] = match['poc']
        else:
            d_obj['gap_target'] = 7
            d_obj['gap_achieved'] = d_obj['filled_posts']
            d_obj['gap_remaining'] = d_obj['vacant_posts']

    zones_list = []
    for z_name, z_data in sorted(zones_map.items()):
        zones_list.append({
            'name': z_name,
            'districts': sorted(list(z_data['districts'])),
            'district_count': len(z_data['districts']),
            'total_posts': z_data['total_posts'],
            'filled_posts': z_data['filled_posts'],
            'vacant_posts': z_data['vacant_posts'],
            'fill_rate': round((z_data['filled_posts'] / z_data['total_posts']) * 100, 1) if z_data['total_posts'] else 0,
            'zonal_lead': z_data['zonal_lead'] or 'Not Appointed'
        })

    districts_list = []
    for d_name, d_data in sorted(districts_map.items()):
        # Calculate Action Priority Tier
        rem = d_data['gap_remaining']
        has_dp = bool(d_data['district_president'])
        has_dgs = bool(d_data['district_gen_sec'])
        
        tier = 'Vanguard' # Complete
        if rem >= 4 or not has_dp or not has_dgs:
            tier = 'Critical'
        elif rem > 0:
            tier = 'Developing'
            
        districts_list.append({
            'name': d_name,
            'zone': d_data['zone'],
            'total_posts': d_data['total_posts'],
            'filled_posts': d_data['filled_posts'],
            'vacant_posts': d_data['vacant_posts'],
            'fill_rate': round((d_data['filled_posts'] / d_data['total_posts']) * 100, 1) if d_data['total_posts'] else 0,
            'district_president': d_data['district_president'] or '(Vacant)',
            'district_gen_sec': d_data['district_gen_sec'] or '(Vacant)',
            'has_dp': has_dp,
            'has_dgs': has_dgs,
            'core_committee_count': d_data['core_committee_count'],
            'gap_target': d_data['gap_target'],
            'gap_achieved': d_data['gap_achieved'],
            'gap_remaining': d_data['gap_remaining'],
            'priority_tier': tier,
            'poc': d_data['poc']
        })

    # -------------------------------------------------------------
    # 6. Deep Analytics & Reasoning Aggregations
    # -------------------------------------------------------------
    filled_members = [m for m in swg_members if not m['is_vacant']]
    all_active_members = filled_members + [m for m in campaign_committee if not m['is_vacant']]
    
    # Category and Caste (from all active members)
    category_counts = Counter(m['category'] for m in all_active_members if m['category'])
    caste_counts = Counter(m['caste'] for m in all_active_members if m['caste'])
    
    # Political Migrations
    political_migration_counts = Counter()
    for m in all_active_members:
        for p in m.get('parties', []):
            political_migration_counts[p] += 1
            
    # Business Sectors
    sector_counts = Counter()
    for m in all_active_members:
        for s in m.get('sectors', []):
            sector_counts[s] += 1
            
    # Age Cohorts
    age_cohorts = {'Under 35 (Youth)': 0, '35-49 (Prime)': 0, '50-59 (Senior)': 0, '60+ (Veterans)': 0, 'Unspecified': 0}
    for m in all_active_members:
        a = m['age']
        if a:
            if a < 35: age_cohorts['Under 35 (Youth)'] += 1
            elif a < 50: age_cohorts['35-49 (Prime)'] += 1
            elif a < 60: age_cohorts['50-59 (Senior)'] += 1
            else: age_cohorts['60+ (Veterans)'] += 1
        else:
            age_cohorts['Unspecified'] += 1
            
    posts_counts = Counter(m['post'] for m in swg_members if m['post'])
    posts_filled = Counter(m['post'] for m in filled_members if m['post'])

    # District Action Priorities summary
    priority_counts = Counter(d['priority_tier'] for d in districts_list)

    demographics = {
        'categories': dict(category_counts.most_common()),
        'castes': dict(caste_counts.most_common(20)),
        'political_migration': dict(political_migration_counts.most_common()),
        'business_sectors': dict(sector_counts.most_common(15)),
        'age_cohorts': age_cohorts,
        'calling_status': dict(calling_status_counts.most_common()),
        'district_priority_tiers': dict(priority_counts),
        'posts_distribution': {
            p: {
                'total': posts_counts[p],
                'filled': posts_filled[p],
                'vacant': posts_counts[p] - posts_filled[p],
                'fill_rate': round((posts_filled[p] / posts_counts[p]) * 100, 1)
            } for p in posts_counts
        }
    }

    import datetime
    now_iso = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # -------------------------------------------------------------
    # 6. Multilingual Hindi Enrichment
    # -------------------------------------------------------------
    try:
        from translations_data import (
            transliterate_name, translate_profile,
            DISTRICTS_HI, ZONES_HI, POSTS_HI, CATEGORIES_HI, CASTES_HI, PRIORITY_TIERS_HI
        )
        for m in campaign_committee:
            if m.get('name'):
                m['name_hi'] = transliterate_name(m['name'])
            if m.get('profile'):
                m['profile_hi'] = translate_profile(m['profile'], m.get('name_hi', ''))
            m['district_hi'] = DISTRICTS_HI.get(m.get('district'), m.get('district'))
            m['zone_hi'] = ZONES_HI.get(m.get('zone'), m.get('zone'))
            m['caste_hi'] = CASTES_HI.get(m.get('caste'), m.get('caste'))
            m['category_hi'] = CATEGORIES_HI.get(m.get('category'), m.get('category'))

        for m in swg_members:
            if m.get('name'):
                m['name_hi'] = transliterate_name(m['name'])
            if m.get('profile'):
                m['profile_hi'] = translate_profile(m['profile'], m.get('name_hi', ''))
            m['district_hi'] = DISTRICTS_HI.get(m.get('district'), m.get('district'))
            m['zone_hi'] = ZONES_HI.get(m.get('zone'), m.get('zone'))
            m['post_hi'] = POSTS_HI.get(m.get('post'), m.get('post'))
            m['caste_hi'] = CASTES_HI.get(m.get('caste'), m.get('caste'))
            m['category_hi'] = CATEGORIES_HI.get(m.get('category'), m.get('category'))

        for d in districts_list:
            d['name_hi'] = DISTRICTS_HI.get(d['name'], d['name'])
            d['zone_hi'] = ZONES_HI.get(d.get('zone'), d.get('zone'))
            if d.get('district_president'):
                d['district_president_hi'] = transliterate_name(d['district_president'])
            if d.get('district_gen_sec'):
                d['district_gen_sec_hi'] = transliterate_name(d['district_gen_sec'])
            if d.get('priority_tier'):
                d['priority_tier_hi'] = PRIORITY_TIERS_HI.get(d['priority_tier'], d['priority_tier'])

        for z in zones_list:
            z['name_hi'] = ZONES_HI.get(z['name'], z['name'])
    except Exception as e:
        print(f"Hindi enrichment warning: {e}")

    full_data = {
        'last_updated': now_iso,
        'source_sheet_url': 'https://docs.google.com/spreadsheets/d/1CifDZWm_rhf6abDpKf-5ZH2hLQnFTmJZoaNssKZjA_g/edit',
        'summary': {
            'total_positions': len(swg_members),
            'filled_positions': len(filled_members),
            'vacant_positions': len(swg_members) - len(filled_members),
            'overall_fill_rate': round((len(filled_members) / len(swg_members)) * 100, 1),
            'campaign_committee_count': len(campaign_committee),
            'zones_count': len(zones_list),
            'districts_count': len(districts_list),
            'grassroots_elected_count': sum(1 for m in filled_members if 'Grassroots Elected' in m['sectors']),
            'zero_gap_districts_count': sum(1 for d in districts_list if d['gap_remaining'] == 0),
            'critical_districts_count': priority_counts['Critical']
        },
        'campaign_committee': campaign_committee,
        'swg_members': swg_members,
        'zones': zones_list,
        'districts': districts_list,
        'demographics': demographics,
        'tirhut_calling': tirhut_calls
    }

    return full_data

if __name__ == '__main__':
    data = process_all_data('data.xlsx')
    with open('dashboard_data.json', 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    with open('data.js', 'w', encoding='utf-8') as f:
        f.write('window.DASHBOARD_DATA = ' + json.dumps(data, ensure_ascii=False) + ';')

    print(f"Successfully processed and updated data! Last updated: {data['last_updated']}")
    print(f"Political migrations: {data['demographics']['political_migration']}")
    print(f"Business sectors: {data['demographics']['business_sectors']}")
    print(f"District tiers: {data['demographics']['district_priority_tiers']}")
